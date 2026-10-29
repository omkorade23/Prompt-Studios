import pptx
import sys

def inspect_ppt(filename):
    prs = pptx.Presentation(filename)
    for i, slide in enumerate(prs.slides):
        print(f"\n--- Slide {i+1} ---")
        for j, shape in enumerate(slide.shapes):
            shape_type = shape.shape_type if hasattr(shape, 'shape_type') else 'Unknown'
            name = shape.name if hasattr(shape, 'name') else 'Unnamed'
            geom = "N/A"
            try:
                geom = f"left={shape.left}, top={shape.top}, width={shape.width}, height={shape.height}"
            except:
                pass
            print(f"  Shape {j}: Type={shape_type}, Name='{name}', Geom: {geom}")
            if shape.has_text_frame:
                text = shape.text_frame.text.replace("\n", " \\n ")
                print(f"    Text: '{text[:100]}...'")

if __name__ == "__main__":
    import os
    default_path = os.path.join(os.path.dirname(__file__), "..", "docs", "presentations", "Prompt_Studio_Mini_Project.pptx")
    filename = sys.argv[1] if len(sys.argv) > 1 else default_path
    inspect_ppt(filename)
