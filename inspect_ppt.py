import pptx
import sys

def inspect_ppt(filename):
    prs = pptx.Presentation(filename)
    for i, slide in enumerate(prs.slides):
        print(f"\n--- Slide {i+1} ---")
        for j, shape in enumerate(slide.shapes):
            shape_type = shape.shape_type if hasattr(shape, 'shape_type') else 'Unknown'
            name = shape.name if hasattr(shape, 'name') else 'Unnamed'
            geom = "N/A"
            try:
                geom = f"left={shape.left}, top={shape.top}, width={shape.width}, height={shape.height}"
            except:
                pass
            print(f"  Shape {j}: Type={shape_type}, Name='{name}', Geom: {geom}")
            if shape.has_text_frame:
                text = shape.text_frame.text.replace("\n", " \\n ")
                print(f"    Text: '{text[:100]}...'")

if __name__ == "__main__":
    if len(sys.argv) > 1:
        inspect_ppt(sys.argv[1])
    else:
        print("Provide filename")
