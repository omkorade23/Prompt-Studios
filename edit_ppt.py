import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def replace_text(shape, new_text):
    if shape.has_text_frame:
        shape.text_frame.clear()
        p = shape.text_frame.paragraphs[0]
        p.text = new_text

def add_box(slide, left, top, width, height, text):
    shape = slide.shapes.add_shape(
        pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, 
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(79, 129, 189)
    shape.text = text
    return shape

def add_arrow(slide, left, top, width, height, is_down=True):
    shape_type = pptx.enum.shapes.MSO_SHAPE.DOWN_ARROW if is_down else pptx.enum.shapes.MSO_SHAPE.RIGHT_ARROW
    shape = slide.shapes.add_shape(
        shape_type, 
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(128, 128, 128)
    return shape

def process_ppt():
    prs = pptx.Presentation("Mini Project PPT Sample.pptx")
    
    # Slide 1: Title
    slide1 = prs.slides[0]
    replace_text(slide1.shapes[0], "Prompt Studio\nAI-Powered Chatbot & Prompt Review Platform")

    # Slide 2: Outline
    slide2 = prs.slides[1]
    replace_text(slide2.shapes[1], "Introduction\nProblem Statement\nObjectives\nLiterature Survey\nSystem Architecture\nChatbot Workflow\nDatabase Design\nPrompt Review Workflow\nTechnologies Used\nResults and Output\nConclusion\nFuture Scope")

    # Slide 3: Introduction
    slide3 = prs.slides[2]
    replace_text(slide3.shapes[1], " Prompt Studio is an intelligent conversational AI platform built to manage and analyze prompts.\n It combines a persistent chat interface with seamless AI interaction.\n The system bridges the gap between casual AI chat and structured prompt evaluation.")

    # Slide 4: Problem Statement
    slide4 = prs.slides[3]
    replace_text(slide4.shapes[0], "Problem Statement")
    for s in list(slide4.shapes):
        if s.shape_type == 13: # Picture
            sp = s._element
            sp.getparent().remove(sp)
    txBox = slide4.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(3))
    txBox.text_frame.text = " Generic chatbot interfaces lack persistent, user-specific conversation history management.\n Users struggle with the lack of structured prompt review and evaluation.\n There is a separation between AI interactions and persistent project data storage."

    # Slide 5: Objectives
    slide5 = prs.slides[4]
    replace_text(slide5.shapes[1], " Develop a centralized, responsive web application for AI interactions.\n Implement persistent database storage for users, conversations, and messages.\n Integrate Gemini API (gemini-2.5-flash) for real-time generative responses.\n Create a dedicated workflow for contextual prompt review and feedback.")

    # Slide 6: Literature Survey (Table)
    slide6 = prs.slides[5]
    replace_text(slide6.shapes[1], "Summary of existing AI conversational models and prompt engineering techniques.")
    table = slide6.shapes[2].table
    table.cell(0, 0).text = "Reference"
    table.cell(0, 1).text = "Topic"
    table.cell(0, 2).text = "Relevance"
    
    table.cell(1, 0).text = "Attention Is All You Need (Vaswani et al.)"
    table.cell(1, 1).text = "Transformer architectures"
    table.cell(1, 2).text = "Foundational for LLMs"
    
    table.cell(2, 0).text = "A Prompt Pattern Catalog (White et al., 2023)"
    table.cell(2, 1).text = "Prompt engineering structures"
    table.cell(2, 2).text = "Guides our prompt review logic"

    # Slide 7: Literature Survey (Image -> Text)
    slide7 = prs.slides[6]
    for s in list(slide7.shapes):
        if s.shape_type == 13:
            sp = s._element
            sp.getparent().remove(sp)
    txBox = slide7.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(3))
    txBox.text_frame.text = "Prompt Studio improves on generic interfaces by allowing targeted prompt evaluation and persistence, moving beyond temporary chats."

    # Slide 8: System Architecture
    slide8 = prs.slides[7]
    replace_text(slide8.shapes[0], "System Architecture")
    replace_text(slide8.shapes[1], "")
    for s in list(slide8.shapes):
        if s.shape_type == 13:
            sp = s._element
            sp.getparent().remove(sp)
    
    add_box(slide8, 3, 1.5, 4, 0.8, "Frontend Chat Interface\n(HTML / CSS / JS)")
    add_arrow(slide8, 4.8, 2.4, 0.4, 0.5)
    add_box(slide8, 3, 3, 4, 0.8, "Spring Boot REST API\n(Java 17 / Auth / Services)")
    add_arrow(slide8, 4.8, 3.9, 0.4, 0.5)
    
    add_box(slide8, 1, 4.5, 3.5, 0.8, "PostgreSQL Database\n(Neon)")
    add_box(slide8, 5.5, 4.5, 3.5, 0.8, "Gemini API\n(gemini-2.5-flash)")

    # Slide 9: Chatbot Workflow
    slide9 = prs.slides[8]
    for s in list(slide9.shapes):
        if s.shape_type == 13:
            sp = s._element
            sp.getparent().remove(sp)
    txBox = slide9.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(1))
    txBox.text_frame.text = "Chatbot Workflow"
    
    add_box(slide9, 3, 1.2, 4, 0.5, "User Sends Message")
    add_arrow(slide9, 4.8, 1.8, 0.4, 0.3)
    add_box(slide9, 3, 2.2, 4, 0.5, "Authenticate & Identify Conversation")
    add_arrow(slide9, 4.8, 2.8, 0.4, 0.3)
    add_box(slide9, 3, 3.2, 4, 0.5, "Retrieve History & Send to Gemini")
    add_arrow(slide9, 4.8, 3.8, 0.4, 0.3)
    add_box(slide9, 3, 4.2, 4, 0.5, "Save Assistant Response & Display")

    # Slide 10: Database Design
    slide10 = prs.slides[9]
    for s in list(slide10.shapes):
        if s.shape_type == 13:
            sp = s._element
            sp.getparent().remove(sp)
    txBox = slide10.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(1))
    txBox.text_frame.text = "Database Design (ER Diagram)"
    
    add_box(slide10, 1, 2.5, 2, 1, "USERS\n- id\n- email\n- name")
    add_arrow(slide10, 3.1, 2.8, 1.3, 0.3, is_down=False)
    add_box(slide10, 4.5, 2.5, 2, 1, "CONVERSATIONS\n- id\n- user_id\n- title")
    add_arrow(slide10, 6.6, 2.8, 1.3, 0.3, is_down=False)
    add_box(slide10, 8, 2.5, 2, 1, "MESSAGES\n- id\n- conv_id\n- role\n- content")

    # Slide 11: Prompt Review Workflow
    slide11 = prs.slides[10]
    for s in list(slide11.shapes):
        if s.shape_type == 13:
            sp = s._element
            sp.getparent().remove(sp)
    txBox = slide11.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(1))
    txBox.text_frame.text = "Prompt Review Workflow"
    
    add_box(slide11, 3.5, 1.5, 3, 0.6, "User requests review")
    add_arrow(slide11, 4.8, 2.2, 0.4, 0.4)
    add_box(slide11, 3.5, 2.7, 3, 0.6, "Prompt provided?")
    
    # Yes branch
    add_box(slide11, 1.5, 4.0, 2.5, 0.6, "Review Prompt (Future Scope)")
    # No branch
    add_box(slide11, 6.0, 4.0, 2.5, 0.6, "Ask user to upload")

    # Slide 12 & 13: Remove or just replace text
    for sl_idx in [11, 12]:
        sl = prs.slides[sl_idx]
        for s in list(sl.shapes):
            if s.shape_type == 13:
                sp = s._element
                sp.getparent().remove(sp)

    # Slide 14: Tech Used
    slide14 = prs.slides[13]
    for s in list(slide14.shapes):
        if s.shape_type == 13:
            sp = s._element
            sp.getparent().remove(sp)
    txBox = slide14.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(4))
    txBox.text_frame.text = "Backend: Java 17, Spring Boot, JPA/Hibernate, Maven\nDatabase: PostgreSQL (Neon)\nAI Integration: Gemini API (gemini-2.5-flash)\nFrontend: HTML, CSS, JavaScript (Vanilla)\nAuth: Google OAuth & Dev Session Auth"

    # Slide 15: Results and Output
    slide15 = prs.slides[14]
    replace_text(slide15.shapes[1], " Chat interface and conversation management successfully implemented.\n Real-time Gemini integration working.\n UI for Prompt Upload isolated and ready for backend hooks.")
    add_box(slide15, 1, 3, 8, 2, "[PASTE SCREENSHOTS OF LOCALHOST:8000 CHAT UI HERE]")

    # Slide 16: Conclusion
    slide16 = prs.slides[15]
    replace_text(slide16.shapes[1], " Prompt Studio successfully implements a full-stack conversational AI platform.\n The Java/Spring Boot backend effectively manages state in a Neon PostgreSQL database.\n The foundation is laid for advanced prompt review and scoring capabilities.")

    # Slide 17: Future Scope
    slide17 = prs.slides[16]
    replace_text(slide17.shapes[0], "Future Scope")
    replace_text(slide17.shapes[1], " Full prompt analysis and scoring engine.\n File/document-based backend parsing for prompts.\n Analytics dashboard for conversation metrics.\n Prompt versioning and history tracking.")

    # Slide 18: References
    slide18 = prs.slides[17]
    replace_text(slide18.shapes[1], " Gemini API Documentation\n Spring Boot & Hibernate Documentation\n White et al. (2023) - A Prompt Pattern Catalog")

    prs.save("Prompt_Studio_Mini_Project.pptx")
    print("Done generating PPT.")

if __name__ == "__main__":
    process_ppt()
